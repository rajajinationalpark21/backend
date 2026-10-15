"""Turns the client's page documents into page blocks for the website.

    python tools/extract_client_docs.py

Reads the .docx files in documentation/wetransfer_photos-for-website_*/ and
writes src/seed/clientDocPages.json, which src/scripts/applyClientDocPages.js
loads into the content document. Requires python-docx.

Each page becomes a list of blocks the frontend's DocContent component renders:
  h2 / h3 / p / note / ul / ol / table / fact / chips / links / cta / contact / gallery
Editor notes in the documents ("[Create Internal Link]", Hindi reminders, the
SEO / footer sections at the end) are not page text: link and booking notes
become real links and buttons, everything after an SEO or footer marker is dropped.
"""
import json
import re
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
SOURCE = next((ROOT / "documentation").glob("wetransfer_photos-for-website_*"))
OUT = Path(__file__).resolve().parents[1] / "src" / "seed" / "clientDocPages.json"

ZONES = "PAGE - Jungle Safari Jeep Safari Zones"
WILDLIFE = "PAGE - Wildlife & Biodiversity"
PAGES = {
    "zones": {
        "chilla-jeep-safari": f"{ZONES}/Jeep safari Jungle safari Chilla.docx",
        "ranipur-gate-jeep-safari": f"{ZONES}/Jeep safari Jungle safari Ranipur.docx",
        "chillawali-jeep-safari": f"{ZONES}/Jeep Safari Jungle safari Chillawali.docx",
        "motichur-jeep-safari": f"{ZONES}/Jeep Safari Jungle safari Motichur.docx",
        "jhilmil-jheel-safari": f"{ZONES}/Jeep safari Jungle safari Jhilmil.docx",
        "gohri-range-safari": f"{ZONES}/Jeep Safari Jungle safari Gohari.docx",
        "chaurasi-kutiya-beatles-ashram": f"{ZONES}/Chaurasi Kutiya – Beatles Ashram.docx",
    },
    "about": "PAGE - Home About Rajaji/Page - Home About Rajaji.docx",
    "tickets": "PAGE - Ticket and Entry/Tickets, Entry Fees & Safari Charges.docx",
    "wildlife": f"{WILDLIFE}/Main page - Wildlife & Biodiversity.docx",
    "birds": f"{WILDLIFE}/Sub-page - Birds of Rajaji.docx",
    "butterflies": f"{WILDLIFE}/Sub-page - Butterflies of Rajaji.docx",
    "flora": f"{WILDLIFE}/Sub-page - Flora of Rajaji.docx",
    "mammals": f"{WILDLIFE}/Sub-page - Mammals of Rajaji.docx",
    "reptiles": f"{WILDLIFE}/Sub-page - Reptiles of Rajaji.docx",
}

# Everything from the first matching line onwards is notes for the web team.
STOP = re.compile(
    r"^(SEO Information|Strong SEO setup|SEO title:|SEO suggestions|One SEO point|One important SEO"
    r"|Primary keyword|Recommended Image Alt Text|Footer me\b)",
    re.I,
)
DEVANAGARI = re.compile(r"\(?[ऀ-ॿ][ऀ-ॿ\s,।]*\)?,?")
LINK_NOTE = re.compile(
    r"\[\s*(Create\s+)?(Internal\s+Link|Booking\s+(Form\s+)?Link)\s*\]|create (link )?call (button|&|and)[^.]*$"
    r"|create link$"
    r"|\(?\s*\**\s*(create\s+)?(link\s+)?(call|whats\s*app)\s*(&|and|/)\s*(call|whats\s*app)\s*butt\w*(\s+link)?\s*\**\s*\)?",
    re.I,
)
PHONE = re.compile(r"(?:\+?91[\s-]?)?\d{5}\s?\d{5}")

INTERNAL = [
    (r"flora\s*&\s*fauna|wildlife\s*&\s*biodiversity", "/wildlife"),
    (r"tickets", "/rajaji-national-park-jeep-safari-tickets-entry-fees-charges"),
    (r"terms|cancellation", "/rajaji-jeep-safari-terms-conditions-cancellation-refund"),
    (r"mammals", "/mammals-of-rajaji-tiger-reserve"),
    (r"birds", "/birds-of-rajaji-national-park"),
    (r"reptiles", "/reptiles-of-rajaji-tiger-reserve"),
    (r"butterfl", "/butterflies-of-rajaji-national-park"),
    (r"flora", "/flora-of-rajaji-national-park"),
]


def internal_link(label):
    for pattern, to in INTERNAL:
        if re.search(pattern, label, re.I):
            return to
    return None


def numbering_formats(document):
    """numId -> numFmt of level 0 ("bullet", "decimal", ...)."""
    try:
        numbering = document.part.numbering_part.element
    except Exception:
        return {}
    abstract = {}
    for node in numbering.findall(qn("w:abstractNum")):
        lvl = node.find(qn("w:lvl"))
        fmt = lvl.find(qn("w:numFmt")) if lvl is not None else None
        abstract[node.get(qn("w:abstractNumId"))] = fmt.get(qn("w:val")) if fmt is not None else "bullet"
    formats = {}
    for node in numbering.findall(qn("w:num")):
        ref = node.find(qn("w:abstractNumId"))
        formats[node.get(qn("w:numId"))] = abstract.get(ref.get(qn("w:val")) if ref is not None else None, "bullet")
    return formats


def run_is_bold(run, style_bold):
    rpr = run.find(qn("w:rPr"))
    b = rpr.find(qn("w:b")) if rpr is not None else None
    if b is None:
        return style_bold
    return b.get(qn("w:val")) not in ("0", "false")


def paragraph_info(element, document, formats):
    para = docx.text.paragraph.Paragraph(element, document)
    style = para.style
    style_bold = bool(style.font.bold) or style.name.startswith("Heading")
    parts, sizes, notes = [], [], []
    for run in element.iter(qn("w:r")):
        text = ""
        for child in run:
            if child.tag == qn("w:t"):
                text += child.text or ""
            elif child.tag == qn("w:br"):
                text += "\n"
            elif child.tag == qn("w:tab"):
                text += " "
        # Hindi notes for the web team are dropped run by run, so they cannot
        # leave empty bold markers behind; they are kept aside as `notes`.
        notes.extend(DEVANAGARI.findall(text))
        text = DEVANAGARI.sub("", text)
        if not text:
            continue
        sz = run.find(f"{qn('w:rPr')}/{qn('w:sz')}")
        if sz is not None and text.strip():
            sizes.append(int(sz.get(qn("w:val"))) / 2)
        parts.append((text, run_is_bold(run, style_bold)))
    visible = [(t, b) for t, b in parts if re.search(r"\w", t)]
    all_bold = bool(visible) and all(b for _, b in visible)
    # Inline bold is kept as **...** only when part of the paragraph is bold.
    if all_bold or not any(b for _, b in visible):
        text = "".join(t for t, _ in parts)
    else:
        segments = []
        for t, b in parts:
            b = b and bool(re.search(r"\w", t)) or (b and segments and segments[-1][1] and not t.strip())
            if segments and segments[-1][1] == b:
                segments[-1][0] += t
            else:
                segments.append([t, b])
        text = ""
        for t, b in segments:
            if not b:
                text += t
                continue
            for i, line in enumerate(t.split("\n")):
                core = line.strip()
                if i:
                    text += "\n"
                if core:
                    text += line[: len(line) - len(line.lstrip())] + f"**{core}**" + line[len(line.rstrip()):]
                else:
                    text += line
    num = element.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
    list_kind = None
    if num is not None:
        list_kind = "ol" if formats.get(num.get(qn("w:val")), "bullet") not in ("bullet", "none") else "ul"
    return {
        "text": text.strip(),
        "style": style.name,
        "bold": all_bold,
        "size": max(sizes) if sizes else 0,
        "list": list_kind,
        "notes": " ".join(notes),
    }


def clean(text):
    text = DEVANAGARI.sub("", text)
    text = text.replace("Moring", "Morning")
    text = LINK_NOTE.sub("", text)
    text = re.sub(r"\(\s*\)|\[\s*\]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"(\*\*)? +([,.;])", r"\1\2", text)
    text = re.sub(r" *\n *", "\n", text)
    return text.strip(" ,\n")


def bold_labels(text):
    """"Gypsy Safari Cost: ₹3,500" -> "**Gypsy Safari Cost:** ₹3,500" on lines without bold."""
    out = []
    for line in text.split("\n"):
        match = re.match(r"^([A-Z][^:*\n]{1,40}):\s+(\S.*)$", line)
        if match and "**" not in line and not re.search(r"https?$", match.group(1)):
            line = f"**{match.group(1)}:** {match.group(2)}"
        out.append(line)
    return "\n".join(out)


def phones_in(text):
    found = []
    for match in PHONE.findall(text):
        digits = re.sub(r"\D", "", match)[-10:]
        pretty = f"+91 {digits[:5]} {digits[5:]}"
        if pretty not in found:
            found.append(pretty)
    return found


def heading_level(info, is_first):
    style, size, bold, text = info["style"], info["size"], info["bold"], info["text"]
    if info["list"] or not bold or len(text) > 140 or (text.endswith(".") and len(text) > 50):
        return None
    if style.startswith("Heading 3") or re.fullmatch(r"By (Road|Train|Air)", text) or re.match(r"^[^\w\s(\[*]", text):
        return "h3"
    if style.startswith("Heading 1") or style.startswith("Heading 2") or size >= 13.5:
        return "h2"
    if style.startswith("Heading 3") or (not text.endswith(".") and len(text) < 90 and not text.endswith(":")):
        return "h3"
    return None


def table_rows(element, document):
    table = docx.table.Table(element, document)
    rows = []
    for row in table.rows:
        cells, seen = [], set()
        for cell in row.cells:
            if id(cell._tc) in seen:  # merged cells repeat
                continue
            seen.add(id(cell._tc))
            cells.append(clean("\n".join(p.text for p in cell.paragraphs)))
        if any(cells):
            rows.append(cells)
    return rows


def extract(path):
    document = docx.Document(path)
    formats = numbering_formats(document)
    blocks, title, kicker = [], None, []
    pending_label, kicker_from_h1 = None, False

    def push(block):
        nonlocal pending_label
        if block["t"] == "p":
            block["text"] = bold_labels(block["text"])
        elif block["t"] in ("ul", "ol"):
            block["items"] = [bold_labels(item) for item in block["items"]]
        if pending_label and block["t"] in ("p",):
            block["text"] = f"**{pending_label}** {block['text']}"
            pending_label = None
        elif pending_label:
            blocks.append({"t": "p", "text": f"**{pending_label}**"})
            pending_label = None
        if block["t"] in ("ul", "ol") and blocks and blocks[-1]["t"] == block["t"]:
            blocks[-1]["items"].extend(block["items"])
            return
        if block["t"] == "links" and blocks and blocks[-1]["t"] == "links":
            blocks[-1]["items"].extend(block["items"])
            return
        if block["t"] == "cta" and blocks and blocks[-1]["t"] == "cta":
            for item in block["items"]:
                if item not in blocks[-1]["items"]:
                    blocks[-1]["items"].append(item)
            return
        blocks.append(block)

    def push_gallery():
        if blocks and blocks[-1]["t"] == "gallery":
            return
        title = None
        # "Birds of Rajaji –" / "Butterflies of Rajaji – Photo Gallery" headings name the gallery.
        if blocks and blocks[-1]["t"] in ("h2", "h3"):
            title = blocks.pop()["text"].rstrip(" –-")
            if "gallery" not in title.lower():
                title += " – Photo Gallery"
        push({"t": "gallery", **({"title": title} if title else {})})

    for element in document.element.body.iterchildren():
        tag = element.tag.split("}")[1]
        if tag == "tbl":
            rows = table_rows(element, document)
            placeholder = len(rows) <= 2 and not any(cell for row in rows for cell in row[1:])
            if not rows or placeholder or rows[0][0].lower().startswith(("add photos", "photos of")):
                push_gallery()
                continue
            push({"t": "table", "rows": rows})
            continue
        if tag != "p":
            continue
        info = paragraph_info(element, document, formats)
        raw = info["text"]
        if not raw:
            continue
        if STOP.match(raw.replace("**", "").strip()):
            break
        had_link_note = bool(LINK_NOTE.search(raw)) or "→" in raw
        text = clean(raw)
        # "Book Jeep Safari Online (create the booking-form link)" inside a sentence becomes a link.
        if re.search(r"बुकिंग|Booking Form Link", raw + info["notes"]) and "→" not in text:
            text = re.sub(r"(\*\*)?(Book Jeep Safari Online)(\*\*)?", r"[\2](/booking)", text, count=1)
        if not text:
            continue

        # Photo placeholders: the page header and gallery stand in for them.
        if re.fullmatch(r"(\[[^\]]*\]\s*)+", text):
            if re.search(r"CHECK .*AVAILAB", text, re.I):
                push({"t": "cta", "items": [{"kind": "booking", "label": "Check Safari Availability"}]})
            elif re.search(r"WHATSAPP|CALL", text, re.I):
                push({"t": "cta", "items": [{"kind": "call", "label": "Call Us"}, {"kind": "whatsapp", "label": "WhatsApp"}]})
            continue
        # "Photos of butterfly | ..." gallery markers
        if re.match(r"^Photos of ", text, re.I):
            push_gallery()
            continue
        # "By Train- The nearest ..." is a labelled fact, not a heading.
        fact = re.match(r"^(By (?:Road|Train|Air)|Nearest [\w ]+?)\s*[-–]\s*(.+)$", text.replace("**", ""))
        if fact and (info["bold"] or info["style"].startswith("Heading")):
            push({"t": "fact", "label": fact.group(1), "text": fact.group(2)})
            continue

        # Contact lines: phone numbers become Call / WhatsApp buttons.
        if re.search(r"\bcall\b|whats\s*app", text, re.I) and phones_in(text):
            first, _, rest = text.partition("\n")
            tail = PHONE.split(first)[-1]
            lead = re.split(r"call", first, flags=re.I)[0]
            lead = re.sub(r"[*📞👉:]", "", lead).strip(" ,")
            if len(re.findall(r"[A-Za-z]{2,}", tail)) >= 3:
                # The numbers sit inside a sentence: keep the sentence, then the buttons.
                push({"t": "p", "text": re.sub(r"^(\**[📞👉]\**\s*)+", "", first).strip()})
            elif re.match(r"^Check (Safari )?Availability\b.*\bBook", lead, re.I):
                push({"t": "cta", "items": [{"kind": "booking", "label": re.sub(r"\s+(or|and)$", "", lead)}]})
            elif re.search(r"\w{3}", lead):
                push({"t": "p", "text": lead})
            phones = phones_in(text)
            push({"t": "contact", "phones": phones})
            rest = "\n".join(line for line in rest.split("\n") if not phones_in(line)).strip()
            if rest:
                push({"t": "p", "text": rest})
            continue
        if text.lower().startswith("email:"):
            push({"t": "contact", "email": text.split(":", 1)[1].strip()})
            continue

        # "Label → description" internal links and booking links.
        if "→" in text or (had_link_note and re.match(r"^Explore:", text)):
            plain = clean(LINK_NOTE.sub("", raw.replace("**", "")))
            label, _, desc = plain.partition("→")
            label = re.sub(r"^(👉\s*)?Explore:\s*", "", label).strip(" []\n")
            desc = desc.strip(" []\n")
            if re.search(r"book|availability", label, re.I):
                push({"t": "cta", "items": [{"kind": "booking", "label": label}]})
                continue
            to = internal_link(label)
            if to:
                push({"t": "links", "items": [{"label": label, "desc": desc, "to": to}]})
                continue

        # "👉 Check Safari Availability & Book Your Ranipur Jeep Safari" is a booking button.
        plain = text.replace("**", "")
        if (not info["list"] and "\n" not in plain and "|" not in plain and len(plain) < 120 and not plain.endswith(".")
                and not phones_in(plain)
                and re.match(r"^(👉\s*)?Check (Safari )?Availability\b.*\bBook", plain, re.I)):
            push({"t": "cta", "items": [{"kind": "booking", "label": plain.lstrip("👉 ").strip()}]})
            continue

        level = None if info["list"] else heading_level(info, title is None)
        # The document's real H1 wins over a bold line typed above it.
        if title and not blocks and info["style"].startswith("Heading 1") and not kicker_from_h1:
            kicker.insert(0, title)
            title, kicker_from_h1 = text, True
            continue
        if title is None:
            if level or info["style"].startswith("Heading"):
                title = text
                continue
        if level and not blocks and title and info["size"] < 13.5 and not info["style"].startswith("Heading"):
            kicker.append(text)
            continue
        if level:
            if title and not blocks and text == title:
                continue
            if text.endswith(":"):
                pending_label = text
                continue
            if text.lower() == "gallery":
                push_gallery()
                continue
            # A heading with nothing under it before the next section heading.
            if blocks and blocks[-1]["t"] == "h3" and level == "h2" or blocks and blocks[-1]["t"] == level:
                blocks.pop()
            push({"t": level, "text": text})
            continue
        if info["bold"] and text.endswith(":") and len(text) < 60:
            pending_label = text
            continue
        if info["list"]:
            item = text
            if "→" in raw:
                continue
            push({"t": info["list"], "items": [item]})
            continue
        if re.match(r"^(Important|Tip|Note)\b\s*:", text) or pending_label and re.search(r"disclaimer", pending_label, re.I):
            label = pending_label or text.split(":", 1)[0] + ":"
            body = text if pending_label else text.split(":", 1)[1].strip()
            pending_label = None
            blocks.append({"t": "note", "label": label.rstrip(":"), "text": body})
            continue
        if text.count(" | ") >= 2 and not phones_in(text):
            push({"t": "chips", "items": [part.strip() for part in text.split("|") if part.strip()]})
            continue
        if info["style"].startswith("Heading") and not info["list"]:
            push({"t": "h3", "text": text})
            continue
        if not blocks and info["bold"] is False and len(text) < 120 and text.startswith("("):
            kicker.append(text.strip("()"))
            continue
        push({"t": "p", "text": text})

    # Headings left without content before a gallery or at the end are dropped.
    while blocks and blocks[-1]["t"] in ("h2", "h3"):
        blocks.pop()
    return {"title": title, "kicker": kicker, "blocks": blocks}


def main():
    pages = {}
    for key, value in PAGES.items():
        if isinstance(value, dict):
            pages[key] = {slug: extract(SOURCE / rel) for slug, rel in value.items()}
        else:
            pages[key] = extract(SOURCE / value)
    OUT.write_text(json.dumps(pages, ensure_ascii=False, indent=1), encoding="utf-8")
    count = sum(len(p["blocks"]) for v in pages.values() for p in (v.values() if "blocks" not in v else [v]))
    print(f"wrote {OUT} ({count} blocks)")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
